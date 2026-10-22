import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

with open("export-products-raw.json", "r", encoding="utf-8") as f:
    products = json.load(f)

CATEGORY = {
    "Ambalaj Makineleri": ("Packaging Machines", "Опаковъчни машини", "آلات التغليف"),
    "Endüstriyel Streç Film": ("Industrial Stretch Film", "Индустриално стреч фолио", "فيلم استرتش صناعي"),
    "Masura ve Paletler": ("Spools & Pallets", "Шпули и палети", "ماسورة وبالتات"),
    "Gıda Streç Film": ("Food Stretch Film", "Хранителен стреч фолио", "فيلم استرتش غذائي"),
    "Balonlu Naylon": ("Bubble Wrap", "Балонено фолио", "نايلون فقاعي"),
    "Pre-Streç Film": ("Pre-Stretch Film", "Пред-стреч фолио", "فيلم ما قبل الشد"),
    "PE Köpük": ("PE Foam", "PE пяна", "فوم بولي إيثيلين"),
    "Karton Kutu ve Kağıt Çantalar": ("Cardboard Boxes & Paper Bags", "Картонени кутии и хартиени торби", "كراتين وأكياس ورقية"),
    "Jumbo Streç Film": ("Jumbo Stretch Film", "Джъмбо стреч фолио", "فيلم استرتش جامبو"),
    "Hotmelt Koli Bandı": ("Hotmelt Packing Tape", "Хотмелт опаковъчна лента", "شريط لاصق هوتميلت"),
    "PP Çember": ("PP Strapping", "PP лента за стягане", "شريط تربيط PP"),
    "Maskeleme Bandı": ("Masking Tape", "Маскираща лента", "شريط لاصق للطلاء"),
    "Akrilik Koli Bandı": ("Acrylic Packing Tape", "Акрилна опаковъчна лента", "شريط لاصق أكريليك"),
}

ATTR_LABEL = {
    "Kalınlık": ("Thickness", "Дебелина", "السُمك"),
    "Uzunluk": ("Length", "Дължина", "الطول"),
    "Genişlik": ("Width", "Широчина", "العرض"),
    "Ağırlık": ("Weight", "Тегло", "الوزن"),
    "Masura Ağırlığı": ("Core weight", "Тегло на шпулата", "وزن الماسورة"),
}

COLOR = {
    "Şeffaf": ("Clear", "Прозрачен", "شفاف"),
    "Siyah": ("Black", "Черен", "أسود"),
    "Beyaz": ("White", "Бял", "أبيض"),
    "Mavi": ("Blue", "Син", "أزرق"),
    "Turuncu": ("Orange", "Оранжев", "برتقالي"),
    "Yeşil": ("Green", "Зелен", "أخضر"),
    "Sarı": ("Yellow", "Жълт", "أصفر"),
    "Kahverengi": ("Brown", "Кафяв", "بني"),
    "Krem": ("Cream", "Крем", "كريمي"),
    "Kırmızı": ("Red", "Червен", "أحمر"),
}

ATTR_ORDER = ["Kalınlık", "Uzunluk", "Genişlik", "Ağırlık", "Masura Ağırlığı"]


def build_title(cat_tr, attrs, lang_idx):
    cat = CATEGORY.get(cat_tr, (cat_tr, cat_tr, cat_tr))[lang_idx]
    parts = []
    for key in ATTR_ORDER:
        if key in attrs:
            label = ATTR_LABEL[key][lang_idx]
            parts.append(f"{label} {attrs[key]}")
    if "Renk" in attrs:
        color_tr = attrs["Renk"]
        color = COLOR.get(color_tr, (color_tr, color_tr, color_tr))[lang_idx]
        parts.append(color)
    if parts:
        return f"{cat} — " + ", ".join(parts)
    return cat


wb = Workbook()

# ---- Sheet 0: instructions (Arabic) ----
ws0 = wb.active
ws0.title = "اقرأ أولاً"
ws0.sheet_view.rightToLeft = True
bold = Font(name="Arial", bold=True, size=13)
normal = Font(name="Arial", size=11)
lines = [
    ("ملف مراجعة الترجمة - PORTSAID", bold),
    ("", normal),
    ("ورقة \"المنتجات\": عمود العنوان التركي (الحالي على الموقع) مرجعي فقط - لا تعدله.", normal),
    ("الأعمدة الصفراء (English / Bulgarian / Arabic) هي مسودة مبنية تلقائيًا من الفئة والمواصفات المخزّنة فعليًا، راجعها وعدّلها.", normal),
    ("تقدر تفتح الملف، تعدّل عمود لغة واحدة بس، تحفظ، وترجع لعمود لغة تانية في جلسة تانية - الأعمدة مستقلة عن بعضها.", normal),
    ("عمود \"ملاحظات\" فاضي لأي تعليق أو سؤال.", normal),
    ("", normal),
    ("ورقة \"نصوص الموقع\": كلمات وعناوين ثابتة في واجهة الموقع (أسماء الفئات، عناوين الفلاتر، كلمات عامة) - مترجمة فعلًا، للمراجعة والتصحيح فقط.", normal),
    ("", normal),
    ("لما تخلص تعديل، ابعت الملف وهوديها على الموقع.", normal),
]
for i, (text, font) in enumerate(lines, start=1):
    c = ws0.cell(row=i, column=1, value=text)
    c.font = font
    c.alignment = Alignment(horizontal="right", wrap_text=True)
ws0.column_dimensions["A"].width = 110

# ---- Sheet 1: Products ----
ws = wb.create_sheet("المنتجات")
headers = [
    "الفئة (TR)", "الكود", "العنوان الحالي (تركي)",
    "السُمك", "الطول", "العرض", "الوزن", "وزن الماسورة", "اللون (TR)",
    "English (draft)", "Bulgarian (draft)", "Arabic (draft)", "ملاحظات",
]
editable_cols = {10, 11, 12}
header_fill_ref = PatternFill("solid", fgColor="D9D9D9")
header_fill_edit = PatternFill("solid", fgColor="FFF2A6")
for col, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = Font(name="Arial", bold=True)
    cell.fill = header_fill_edit if col in editable_cols else header_fill_ref

row = 2
for p in products:
    attrs = p.get("attrs") or {}
    attrs = {k: v for k, v in attrs.items() if k != "Default option"}
    cat_tr = p.get("category") or ""
    values = [
        cat_tr,
        p.get("handle", ""),
        p.get("title", ""),
        attrs.get("Kalınlık", ""),
        attrs.get("Uzunluk", ""),
        attrs.get("Genişlik", ""),
        attrs.get("Ağırlık", ""),
        attrs.get("Masura Ağırlığı", ""),
        attrs.get("Renk", ""),
        build_title(cat_tr, attrs, 0),
        build_title(cat_tr, attrs, 1),
        build_title(cat_tr, attrs, 2),
        "",
    ]
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.font = Font(name="Arial", size=10)
        if col in editable_cols:
            cell.fill = PatternFill("solid", fgColor="FFFBE0")
    row += 1

widths = [20, 28, 42, 10, 10, 10, 10, 14, 12, 45, 45, 45, 25]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# ---- Sheet 2: Site UI text ----
ws2 = wb.create_sheet("نصوص الموقع")
site_headers = ["المفتاح (تركي)", "English", "Bulgarian", "Arabic"]
for col, h in enumerate(site_headers, start=1):
    cell = ws2.cell(row=1, column=col, value=h)
    cell.font = Font(name="Arial", bold=True)
    cell.fill = header_fill_edit if col > 1 else header_fill_ref

SITE_TEXT = [
    ("BRAND", "BRAND", "МАРКА", "العلامة التجارية"),
    ("Ambalaj Makineleri", "Packaging Machines", "Опаковъчни машини", "آلات التغليف"),
    ("Endüstriyel Streç Film", "Industrial Stretch Film", "Индустриално стреч фолио", "فيلم استرتش صناعي"),
    ("Masura ve Paletler", "Spools & Pallets", "Шпули и палети", "ماسورة وبالتات"),
    ("Gıda Streç Film", "Food Stretch Film", "Хранителен стреч фолио", "فيلم استرتش غذائي"),
    ("Balonlu Naylon", "Bubble Wrap", "Балонено фолио", "نايلون فقاعي"),
    ("Pre-Streç Film", "Pre-Stretch Film", "Пред-стреч фолио", "فيلم ما قبل الشد"),
    ("PE Köpük", "PE Foam", "PE пяна", "فوم بولي إيثيلين"),
    ("Karton Kutu ve Kağıt Çantalar", "Cardboard Boxes & Paper Bags", "Картонени кутии и хартиени торби", "كراتين وأكياس ورقية"),
    ("Jumbo Streç Film", "Jumbo Stretch Film", "Джъмбо стреч фолио", "فيلم استرتش جامبو"),
    ("Hotmelt Koli Bandı", "Hotmelt Packing Tape", "Хотмелт опаковъчна лента", "شريط لاصق هوتميلت"),
    ("PP Çember", "PP Strapping", "PP лента за стягане", "شريط تربيط PP"),
    ("Maskeleme Bandı", "Masking Tape", "Маскираща лента", "شريط لاصق للطلاء"),
    ("Akrilik Koli Bandı", "Acrylic Packing Tape", "Акрилна опаковъчна лента", "شريط لاصق أكريليك"),
    ("Kalınlık", "Thickness", "Дебелина", "السُمك"),
    ("Uzunluk", "Length", "Дължина", "الطول"),
    ("Genişlik", "Width", "Широчина", "العرض"),
    ("Ağırlık", "Weight", "Тегло", "الوزن"),
    ("Masura Ağırlığı", "Core Weight", "Тегло на шпулата", "وزن الماسورة"),
    ("Renk", "Color", "Цвят", "اللون"),
    ("All products", "All products", "Всички продукти", "جميع المنتجات"),
    ("Products", "Products", "Продукти", "المنتجات"),
    ("Categories", "Categories", "Категории", "الفئات"),
    ("SKU", "SKU", "Артикулен номер", "رمز المنتج"),
    ("Price", "Price", "Цена", "السعر"),
    ("Quantity", "Quantity", "Количество", "الكمية"),
    ("Sort by", "Sort by", "Сортиране", "ترتيب حسب"),
    ("Not specified", "Not specified", "Не е посочено", "غير محدد"),
    ("Price on request", "Price on request", "Поискайте оферта за цена", "اطلب عرض سعر"),
    ("Out of stock", "Out of stock", "Изчерпан", "غير متوفر"),
    ("Availability on request", "Availability on request", "Наличност при запитване", "التوفر عند الطلب"),
]
for r, (tr, en, bg, ar) in enumerate(SITE_TEXT, start=2):
    for col, v in enumerate([tr, en, bg, ar], start=1):
        cell = ws2.cell(row=r, column=col, value=v)
        cell.font = Font(name="Arial", size=10)
        if col > 1:
            cell.fill = PatternFill("solid", fgColor="FFFBE0")

for i, w in enumerate([32, 32, 32, 32], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"

wb.save("translation-review.xlsx")
print("done", len(products), "products")
